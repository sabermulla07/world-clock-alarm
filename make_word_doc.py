import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_report():
    doc = Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    navy_color = RGBColor(26, 54, 93)     # #1a365d
    accent_blue = RGBColor(37, 99, 235)   # #2563eb
    dark_gray = RGBColor(51, 65, 85)      # #334155
    code_color = RGBColor(15, 23, 42)     # #0f172a

    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("LABORATORY ASSIGNMENT REPORT")
    title_run.font.name = "Arial"
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title_run.font.color.rgb = navy_color
    title_p.paragraph_format.space_after = Pt(2)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("Department of Computer Engineering / Information Technology")
    sub_run.font.name = "Arial"
    sub_run.font.size = Pt(10)
    sub_run.font.italic = True
    sub_run.font.color.rgb = dark_gray
    sub_p.paragraph_format.space_after = Pt(14)

    # Student Details Table (Clean structured border box with blank fields as requested)
    details_table = doc.add_table(rows=7, cols=2)
    details_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    details_table.autofit = False

    labels = [
        ("Subject Name:", ""),
        ("Assignment No.:", "02"),
        ("Name of the Student:", ""),
        ("Division:", ""),
        ("Batch:", ""),
        ("Roll No:", ""),
        ("PRN Number:", "")
    ]

    for i, (lbl, val) in enumerate(labels):
        row = details_table.rows[i]
        
        # Label cell
        cell_lbl = row.cells[0]
        cell_lbl.width = Inches(2.2)
        p_lbl = cell_lbl.paragraphs[0]
        p_lbl.paragraph_format.space_before = Pt(3)
        p_lbl.paragraph_format.space_after = Pt(3)
        r_lbl = p_lbl.add_run(lbl)
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(10)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = navy_color
        set_cell_background(cell_lbl, "F1F5F9")
        set_cell_margins(cell_lbl, top=80, bottom=80, left=120, right=120)

        # Value cell (Left blank as requested, with subtle dotted baseline or clean space)
        cell_val = row.cells[1]
        cell_val.width = Inches(4.6)
        p_val = cell_val.paragraphs[0]
        p_val.paragraph_format.space_before = Pt(3)
        p_val.paragraph_format.space_after = Pt(3)
        r_val = p_val.add_run(val if val else "")
        r_val.font.name = "Arial"
        r_val.font.size = Pt(10)
        r_val.font.bold = (val == "02")
        r_val.font.color.rgb = dark_gray
        set_cell_margins(cell_val, top=80, bottom=80, left=120, right=120)

    # Spacing
    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    def add_section_header(title_text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(title_text)
        run.font.name = "Arial"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = navy_color
        return p

    # ================= 1. AIM =================
    add_section_header("AIM:")
    p_aim = doc.add_paragraph()
    p_aim.paragraph_format.space_after = Pt(8)
    p_aim.paragraph_format.line_spacing = 1.15
    run_aim = p_aim.add_run(
        "Build an advanced real-time clock dashboard application using React that combines "
        "Analog Clock, Digital Clock, Time zone Management, Alarm Features, and Dynamic UI "
        "Controls into a single responsive system. The application should simulate a professional "
        "world-clock dashboard capable of displaying live time updates across multiple regions "
        "while maintaining smooth animations and optimized performance."
    )
    run_aim.font.name = "Arial"
    run_aim.font.size = Pt(10.5)
    run_aim.font.color.rgb = dark_gray

    # ================= 2. OBJECTIVES =================
    add_section_header("OBJECTIVES:")
    objectives = [
        "To design and implement a synchronized dual-clock system featuring high-precision vector SVG analog clock faces and a digital clock interface with continuous sweep and discrete quartz movement options.",
        "To develop a multi-region global timezone management network with synchronized mini-clocks, daylight/night tracking, business hour classification, and an interactive 24-hour timeline meeting planner.",
        "To implement an offline Web Audio API synthesized alarm sound engine with multi-tone frequency modulation, repeat schedules, customizable snooze capabilities, and browser push notification integration.",
        "To architect dynamic UI controls and state management using custom React hooks and LocalStorage persistence for high-frequency drift-corrected rendering, theme customization, and responsive layout scaling."
    ]
    for obj in objectives:
        p_obj = doc.add_paragraph(style='List Bullet')
        p_obj.paragraph_format.space_after = Pt(4)
        p_obj.paragraph_format.line_spacing = 1.15
        r_obj = p_obj.add_run(obj)
        r_obj.font.name = "Arial"
        r_obj.font.size = Pt(10)
        r_obj.font.color.rgb = dark_gray

    # ================= 3. SOFTWARE REQUIREMENTS =================
    add_section_header("SOFTWARE REQUIREMENTS:")
    
    req_table = doc.add_table(rows=7, cols=2)
    req_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    req_table.autofit = False

    req_data = [
        ("Operating System", "Windows 10 / 11, Linux, or macOS"),
        ("Runtime Environment", "Node.js (v18.0.0 or higher), npm (v9.0.0 or higher)"),
        ("Frontend Library / Framework", "React 18.3.1, React DOM 18.3.1"),
        ("Build Tool / Bundler", "Vite 6.0.5, @vitejs/plugin-react"),
        ("Styling & UI Tools", "Tailwind CSS 3.4.17, PostCSS, Autoprefixer, Lucide React, Canvas Confetti"),
        ("Web APIs Leveraged", "Web Audio API (Synthesizer), requestAnimationFrame, Fullscreen API, Web Notifications API, LocalStorage API, Intl.DateTimeFormat"),
        ("Code Editor & Browser", "Visual Studio Code, Google Chrome / Microsoft Edge (Latest)")
    ]

    for i, (param, spec) in enumerate(req_data):
        row = req_table.rows[i]
        
        c0 = row.cells[0]
        c0.width = Inches(2.2)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(2)
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(param)
        r0.font.name = "Arial"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = navy_color
        set_cell_background(c0, "F8FAFC")
        set_cell_margins(c0, top=60, bottom=60, left=100, right=100)

        c1 = row.cells[1]
        c1.width = Inches(4.6)
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(spec)
        r1.font.name = "Arial"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = dark_gray
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ================= 4. CODE (FILE STRUCTURE) =================
    add_section_header("CODE:")
    
    p_code_intro = doc.add_paragraph()
    p_code_intro.paragraph_format.space_after = Pt(4)
    r_code_intro = p_code_intro.add_run("Project Architecture & Modular File Structure:")
    r_code_intro.font.name = "Arial"
    r_code_intro.font.size = Pt(10)
    r_code_intro.font.bold = True
    r_code_intro.font.color.rgb = dark_gray

    # Code Box / File Structure Tree
    file_structure_text = """LAB 02/
├── index.html                     # Application HTML entry shell & typography fonts
├── package.json                   # Dependencies, build scripts & metadata
├── vite.config.js                 # Vite compilation & React plugin configuration
├── tailwind.config.js             # Tailwind design tokens, fonts & keyframe animations
├── postcss.config.js              # PostCSS autoprefixer pipeline
├── README.md                      # Complete system documentation & usage guide
└── src/
    ├── main.jsx                   # Application bootstrap & DOM root renderer
    ├── index.css                  # Global Tailwind directives & dark/light theme definitions
    ├── App.jsx                    # Central state manager, tab navigator & layout orchestrator
    ├── constants/
    │   └── timezones.js           # Worldwide timezone database (20+ cities) & offset helpers
    ├── utils/
    │   └── soundEffects.js        # Web Audio API procedural synthesizer (alarms & clock ticks)
    ├── hooks/
    │   ├── useRealTime.js         # Drift-compensated timer with requestAnimationFrame (60 FPS)
    │   ├── useAlarms.js           # Alarm scheduler, notification handler & snooze engine
    │   └── useLocalStorage.js    # Persistent browser storage hook for settings & alarms
    └── components/
        ├── Header.jsx             # Top bar with live status, theme switcher & fullscreen toggle
        ├── AnalogClock.jsx        # Vector SVG clock with 4 face styles & continuous sweep
        ├── DigitalClock.jsx       # Large digital readout with calendar, Day of Year & UTC pill
        ├── HeroClockSection.jsx   # Synchronized dual analog & digital centerpiece view
        ├── WorldClockCard.jsx     # Pinned city card with mini analog face & business hour tags
        ├── WorldClockSection.jsx  # Multi-region world clock network with searchable city adder
        ├── TimeConverter.jsx      # Interactive 24-hour time scrubber & global meeting planner
        ├── AlarmSection.jsx       # Alarm manager with sound tone previews & repeat selectors
        ├── AlarmModal.jsx         # Fullscreen alarm ringing alert with confetti & audio loop
        ├── StopwatchTimerSection.jsx # Millisecond lap stopwatch & Pomodoro countdown timer
        └── SettingsModal.jsx      # Dynamic UI controls & theme customization modal"""

    # Add code box table
    code_table = doc.add_table(rows=1, cols=1)
    code_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    code_table.autofit = False
    
    cell_code = code_table.rows[0].cells[0]
    cell_code.width = Inches(6.8)
    set_cell_background(cell_code, "0F172A") # Dark slate background for code
    set_cell_margins(cell_code, top=140, bottom=140, left=180, right=180)

    p_c = cell_code.paragraphs[0]
    p_c.paragraph_format.line_spacing = 1.05
    p_c.paragraph_format.space_before = Pt(0)
    p_c.paragraph_format.space_after = Pt(0)
    
    r_c = p_c.add_run(file_structure_text)
    r_c.font.name = "Consolas"
    r_c.font.size = Pt(8.5)
    r_c.font.color.rgb = RGBColor(226, 232, 240) # Light slate font

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Component summary table
    p_table_title = doc.add_paragraph()
    p_table_title.paragraph_format.space_before = Pt(6)
    p_table_title.paragraph_format.space_after = Pt(4)
    r_tt = p_table_title.add_run("Module Descriptions & Key Responsibilities:")
    r_tt.font.name = "Arial"
    r_tt.font.size = Pt(10)
    r_tt.font.bold = True
    r_tt.font.color.rgb = dark_gray

    mod_table = doc.add_table(rows=8, cols=2)
    mod_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    mod_table.autofit = False

    mod_data = [
        ("AnalogClock.jsx", "Renders SVG vector dials with 4 face styles (Modern, Roman, Chrono, Minimal) with smooth 60 FPS sweep hand."),
        ("DigitalClock.jsx", "Displays formatted 12H/24H time, milliseconds, weekday, date, Day of Year, Week of Year, and UTC offset."),
        ("WorldClockSection.jsx", "Synchronized world clock grid across 20+ cities with live mini clocks, relative time differences, and business hours."),
        ("TimeConverter.jsx", "Interactive 24-hour timeline scrubber slider calculating converted times across all pinned world cities simultaneously."),
        ("soundEffects.js", "Pure Web Audio API procedural sound synthesizer for 5 distinct alarm melodies and subtle mechanical clock ticks."),
        ("useAlarms.js & AlarmSection.jsx", "Alarm scheduling engine supporting repeat rules, snooze logic, desktop push notifications, and ringing overlays."),
        ("StopwatchTimerSection.jsx", "Precision stopwatch with lap recording (fastest/slowest lap highlighting) and countdown timer with Pomodoro presets."),
        ("SettingsModal.jsx & Header.jsx", "Dynamic UI controls for 4 color themes (Obsidian, Neon, Gold, Light), tick audio toggle, and fullscreen mode.")
    ]

    for i, (m_file, m_desc) in enumerate(mod_data):
        row = mod_table.rows[i]
        
        c0 = row.cells[0]
        c0.width = Inches(2.2)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(2)
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(m_file)
        r0.font.name = "Consolas"
        r0.font.size = Pt(9)
        r0.font.bold = True
        r0.font.color.rgb = accent_blue
        set_cell_background(c0, "F8FAFC")
        set_cell_margins(c0, top=50, bottom=50, left=80, right=80)

        c1 = row.cells[1]
        c1.width = Inches(4.6)
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(m_desc)
        r1.font.name = "Arial"
        r1.font.size = Pt(9)
        r1.font.color.rgb = dark_gray
        set_cell_margins(c1, top=50, bottom=50, left=80, right=80)

    # ================= 5. OUTPUT =================
    add_section_header("OUTPUT:")

    p_out_intro = doc.add_paragraph()
    p_out_intro.paragraph_format.space_after = Pt(6)
    r_out_intro = p_out_intro.add_run(
        "Application is deployed and running live at: http://localhost:5173/\n"
        "Below are the primary dashboard output views and screenshot placeholders:"
    )
    r_out_intro.font.name = "Arial"
    r_out_intro.font.size = Pt(10)
    r_out_intro.font.color.rgb = dark_gray

    output_views = [
        ("Figure 1: Synchronized Dual Clock Dashboard (Analog & Digital Hero View)",
         "Features vector SVG analog clock with smooth continuous sweep hand, modern face dial, sun/moon solar indicator, accompanied by digital clock displaying 12H/24H format, day of year, week number, and UTC offset."),
        ("Figure 2: World Clock Network (Multi-Region Timezone Grid)",
         "Displays synchronized live mini analog and digital clocks across pinned world cities (London, New York, Tokyo, Mumbai, Sydney) with relative offset badges and live business hour indicators."),
        ("Figure 3: Interactive Time Zone Converter & Meeting Planner",
         "Interactive 24-hour timeline scrubber slider allowing dynamic time exploration forwards and backwards across global cities simultaneously."),
        ("Figure 4: Smart Alarms & Web Audio Synthesizer Interface",
         "Alarm manager with repeat schedules, active toggles, sound tone preview selectors, and fullscreen ringing alert modal with snooze and dismiss controls."),
        ("Figure 5: Precision Stopwatch & Countdown Timer Utilities",
         "Millisecond lap stopwatch highlighting fastest lap in emerald green and slowest in red, alongside customizable countdown timer with animated progress bar."),
        ("Figure 6: Dynamic UI Controls & Theme Customization Modal",
         "Settings modal showcasing color theme palette (Obsidian Slate, Cyberpunk Neon, Imperial Gold, Clean Frost), clock face styles, smooth sweep toggle, and mechanical tick sound controls.")
    ]

    for title, desc in output_views:
        # Output box
        box_table = doc.add_table(rows=1, cols=1)
        box_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        box_table.autofit = False

        c_box = box_table.rows[0].cells[0]
        c_box.width = Inches(6.8)
        set_cell_background(c_box, "F8FAFC")
        set_cell_margins(c_box, top=140, bottom=140, left=180, right=180)

        p_b = c_box.paragraphs[0]
        p_b.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_b.paragraph_format.space_before = Pt(20)
        p_b.paragraph_format.space_after = Pt(20)

        r_box_text = p_b.add_run("[ PASTE OUTPUT SCREENSHOT HERE ]\n")
        r_box_text.font.name = "Arial"
        r_box_text.font.size = Pt(11)
        r_box_text.font.bold = True
        r_box_text.font.color.rgb = RGBColor(148, 163, 184) # Muted slate

        r_sub = p_b.add_run("(Press Win + Shift + S to capture from http://localhost:5173/ and Ctrl + V to paste)")
        r_sub.font.name = "Arial"
        r_sub.font.size = Pt(8.5)
        r_sub.font.italic = True
        r_sub.font.color.rgb = RGBColor(148, 163, 184)

        # Caption
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(4)
        p_cap.paragraph_format.space_after = Pt(2)
        r_cap = p_cap.add_run(title)
        r_cap.font.name = "Arial"
        r_cap.font.size = Pt(9.5)
        r_cap.font.bold = True
        r_cap.font.color.rgb = navy_color

        # Description
        p_d = doc.add_paragraph()
        p_d.paragraph_format.space_before = Pt(0)
        p_d.paragraph_format.space_after = Pt(12)
        r_d = p_d.add_run(desc)
        r_d.font.name = "Arial"
        r_d.font.size = Pt(9)
        r_d.font.color.rgb = dark_gray

    # Save document
    output_filename = "WT_Assignment_02_Report.docx"
    doc.save(output_filename)
    print(f"Document saved successfully as {output_filename}")

if __name__ == "__main__":
    create_report()
